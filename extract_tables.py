#!/usr/bin/env python3
"""
Extract PRINTED table data from 53 images of Vietnamese education documents.
Filters out handwritten annotations (pen marks in red/blue/black ink).
Images are organized by class: Lớp 10 (1-16), Lớp 11 (17-35), Lớp 12 (36-53)
Table structure: 9 columns
"""

import pytesseract
from PIL import Image, ImageEnhance, ImageFilter
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import cv2
import os
import re

IMG_DIR = os.path.dirname(os.path.abspath(__file__))

# Class boundaries (determined from OCR scan)
CLASS_RANGES = {
    "Lớp 10": (1, 16),
    "Lớp 11": (17, 35),
    "Lớp 12": (36, 53),
}

# Table column headers
HEADERS = [
    "STT",
    "Năng lực thành phần",
    "Biểu hiện năng lực thành phần",
    "Yêu cầu cần đạt",
    "Mạch nội dung",
    "Đơn vị kiến thức",
    "Các dạng câu hỏi",
    "Mức độ",
    "Hình thức câu hỏi",
]

# OCR confidence threshold: higher = stricter filtering of handwritten text
# Printed text typically scores 60+, handwritten scores below 40
CONFIDENCE_THRESHOLD = 40

# Maximum text truncation for Excel cells
MAX_CELL_TEXT = 5000
MAX_RAW_TEXT = 32000


def remove_colored_ink(image_path):
    """Remove colored (red/blue) pen marks while preserving printed black text."""
    img = cv2.imread(image_path)
    if img is None:
        return None
    hsv = cv2.cvtColor(img, cv2.COLOR_BGR2HSV)

    # Red ink in HSV space (two ranges since red wraps around 0/180)
    mask_red1 = cv2.inRange(hsv, np.array([0, 70, 70]), np.array([15, 255, 255]))
    mask_red2 = cv2.inRange(hsv, np.array([165, 70, 70]), np.array([180, 255, 255]))
    # Blue ink
    mask_blue = cv2.inRange(hsv, np.array([95, 70, 70]), np.array([140, 255, 255]))

    mask_ink = mask_red1 | mask_red2 | mask_blue

    # Dilate to cover ink edges/bleed
    kernel = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (5, 5))
    mask_ink = cv2.dilate(mask_ink, kernel, iterations=1)

    # Replace colored ink with white
    cleaned = img.copy()
    cleaned[mask_ink > 0] = [255, 255, 255]
    return cleaned


def preprocess_for_ocr(cv_img, contrast=1.5):
    """Convert OpenCV image to enhanced PIL image for OCR."""
    rgb = cv2.cvtColor(cv_img, cv2.COLOR_BGR2RGB)
    pil_img = Image.fromarray(rgb)
    gray = pil_img.convert('L')
    enhancer = ImageEnhance.Contrast(gray)
    enhanced = enhancer.enhance(contrast)
    enhanced = enhanced.filter(ImageFilter.SHARPEN)
    return enhanced


def is_garbage(text):
    """Return True if text is likely OCR noise or handwriting artifact."""
    t = text.strip()
    if not t:
        return True
    # Mostly non-alphabetic characters
    alpha = sum(1 for c in t if c.isalpha())
    if len(t) > 3 and alpha / len(t) < 0.25:
        return True
    # Pure punctuation/symbols
    if re.match(r'^[^a-zA-ZÀ-ỹ0-9]+$', t):
        return True
    # Consecutive uppercase nonsense (e.g. "NENGGEGSGPcnv")
    if re.search(r'[A-Z]{6,}', t) and not any(
        w in t for w in ['MORPH', 'THRESH', 'BGR', 'HSV']
    ):
        return True
    return False


def clean_text(text):
    """Clean OCR text: remove artifacts, normalize whitespace."""
    if not text:
        return ""
    # Remove stray pipe/backslash characters (table line artifacts)
    text = re.sub(r'[|\\]', ' ', text)
    # Remove stray single special chars between spaces
    text = re.sub(r'\s[^\w\s]\s', ' ', text)
    # Collapse whitespace
    text = re.sub(r'\s+', ' ', text).strip()
    # Remove leading/trailing dashes or dots
    text = text.strip('-–—.·,;: ')
    return text


def get_word_data(cv_img, w, h):
    """
    Run Tesseract on a cleaned image, return word-level data.
    Tries multiple contrast levels and picks the one with most words.
    """
    best_tsv = pd.DataFrame()

    for contrast in [1.5, 1.2, 1.0, 2.0]:
        enhanced = preprocess_for_ocr(cv_img, contrast=contrast)
        tsv = pytesseract.image_to_data(
            enhanced, lang='vie', output_type=pytesseract.Output.DATAFRAME
        )
        tsv['text'] = tsv['text'].astype(str)
        tsv = tsv[tsv['text'].str.strip() != '']
        tsv = tsv[tsv['text'] != 'nan']
        # Use confidence threshold to filter handwritten text
        tsv = tsv[tsv['conf'] > CONFIDENCE_THRESHOLD]

        if len(tsv) > len(best_tsv):
            best_tsv = tsv

    tsv = best_tsv
    if len(tsv) == 0:
        return pd.DataFrame()

    tsv['cx'] = tsv['left'] + tsv['width'] / 2
    tsv['cy'] = tsv['top'] + tsv['height'] / 2
    tsv['text'] = tsv['text'].str.strip()

    # Filter out garbage words
    tsv = tsv[~tsv['text'].apply(is_garbage)]

    return tsv


def get_raw_printed_text(cv_img):
    """Get raw OCR text from cleaned image, picking the best result."""
    best = ""
    best_word_count = 0

    # Try PSM 3 (auto) and PSM 4 (column) — PSM 6 gives too much noise on photos
    for psm in [3, 4]:
        for contrast in [1.5, 1.2, 1.0]:
            enhanced = preprocess_for_ocr(cv_img, contrast=contrast)
            text = pytesseract.image_to_string(
                enhanced, lang='vie', config=f'--psm {psm}'
            )
            # Count real words (3+ chars, mostly alphabetical)
            words = [
                w for w in text.split()
                if len(w) >= 3 and sum(c.isalpha() for c in w) / len(w) > 0.5
            ]
            if len(words) > best_word_count:
                best = text
                best_word_count = len(words)

    # Clean line by line
    lines = best.split('\n')
    cleaned_lines = []
    for line in lines:
        line = clean_text(line)
        if line and not is_garbage(line):
            cleaned_lines.append(line)
    return '\n'.join(cleaned_lines)


def detect_column_boundaries(cv_img):
    """Detect vertical table lines to find column boundaries."""
    gray = cv2.cvtColor(cv_img, cv2.COLOR_BGR2GRAY)
    h, w = gray.shape

    _, binary = cv2.threshold(gray, 160, 255, cv2.THRESH_BINARY_INV)

    v_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (2, h // 8))
    v_lines = cv2.morphologyEx(binary, cv2.MORPH_OPEN, v_kernel, iterations=2)

    col_sum = np.sum(v_lines, axis=0) / 255
    threshold = h * 0.15
    line_positions = np.where(col_sum > threshold)[0]

    if len(line_positions) == 0:
        return None

    groups = []
    current_group = [line_positions[0]]
    for pos in line_positions[1:]:
        if pos - current_group[-1] < 30:
            current_group.append(pos)
        else:
            groups.append(int(np.mean(current_group)))
            current_group = [pos]
    groups.append(int(np.mean(current_group)))

    return groups


def assign_to_columns(tsv, col_boundaries, w):
    """Assign each word to a column based on x-position."""
    if col_boundaries is None or len(col_boundaries) < 2:
        return None

    if col_boundaries[0] > w * 0.05:
        col_boundaries = [0] + col_boundaries
    if col_boundaries[-1] < w * 0.95:
        col_boundaries = col_boundaries + [w]

    def get_col(cx):
        for i in range(len(col_boundaries) - 1):
            if col_boundaries[i] <= cx < col_boundaries[i + 1]:
                return i
        return len(col_boundaries) - 2

    tsv = tsv.copy()
    tsv['col_idx'] = tsv['cx'].apply(get_col)
    return tsv


def group_into_rows(tsv, threshold=45):
    """Group words into rows based on y-position."""
    if len(tsv) == 0:
        return []

    tsv_sorted = tsv.sort_values('top')
    rows = []
    current_y = tsv_sorted.iloc[0]['top']
    current_indices = [tsv_sorted.index[0]]

    for idx in tsv_sorted.index[1:]:
        if abs(tsv_sorted.loc[idx, 'top'] - current_y) < threshold:
            current_indices.append(idx)
        else:
            rows.append(current_indices)
            current_indices = [idx]
            current_y = tsv_sorted.loc[idx, 'top']
    rows.append(current_indices)
    return rows


def merge_row_cells(tsv, row_indices, num_cols):
    """Merge words in a row into per-column text strings."""
    row = tsv.loc[row_indices]
    cells = {}
    for ci in range(num_cols):
        words = row[row['col_idx'] == ci].sort_values('left')
        text = ' '.join(str(t) for t in words['text'].tolist())
        cells[ci] = clean_text(text)
    return cells


def extract_table(image_path):
    """
    Main extraction: clean image, detect table, OCR cells.
    Returns (table_rows, raw_text).
    """
    # Step 1: Remove colored ink (handwriting)
    cv_img = remove_colored_ink(image_path)
    if cv_img is None:
        return [], ""

    h, w = cv_img.shape[:2]

    # Step 2: Get raw text for reference
    raw_text = get_raw_printed_text(cv_img)

    # Step 3: Get word-level OCR data
    tsv = get_word_data(cv_img, w, h)
    if len(tsv) == 0:
        return [], raw_text

    tsv['text'] = tsv['text'].astype(str)

    # Step 4: Detect column boundaries
    col_bounds = detect_column_boundaries(cv_img)

    if col_bounds is not None and len(col_bounds) >= 8:
        tsv = assign_to_columns(tsv, col_bounds, w)
    else:
        # Fallback: approximate column positions for the 9-column table
        fracs = [0.0, 0.04, 0.10, 0.17, 0.38, 0.49, 0.62, 0.80, 0.88, 1.0]
        col_bounds = [int(f * w) for f in fracs]
        tsv = assign_to_columns(tsv, col_bounds, w)

    if tsv is None:
        return [], raw_text

    # Step 5: Group into rows and merge
    rows = group_into_rows(tsv)
    num_cols = max(tsv['col_idx'].max() + 1, 9)

    table_rows = []
    for row_indices in rows:
        cells = merge_row_cells(tsv, row_indices, num_cols)
        row_list = [cells.get(i, '') for i in range(num_cols)]
        # Keep row only if it has meaningful content
        meaningful = [r for r in row_list if r and len(r) > 1]
        if meaningful:
            table_rows.append(row_list)

    return table_rows, raw_text


def process_all_images():
    """Process all 53 images, organized by class."""
    all_data = {}

    for class_name, (start, end) in CLASS_RANGES.items():
        print(f"\n{'='*60}")
        print(f"Processing {class_name} (images {start}-{end})")
        print(f"{'='*60}")

        class_data = []
        for i in range(start, end + 1):
            path = os.path.join(IMG_DIR, f"{i}.jpg")
            if not os.path.exists(path):
                print(f"  ⚠ {i}.jpg not found, skipping")
                continue

            print(f"  Processing {i}.jpg ...", end=" ", flush=True)
            try:
                table_rows, raw_text = extract_table(path)
                class_data.append({
                    'image': i,
                    'table_rows': table_rows,
                    'raw_text': raw_text,
                })
                print(f"→ {len(table_rows)} rows")
            except Exception as e:
                print(f"ERROR: {e}")
                class_data.append({
                    'image': i,
                    'table_rows': [],
                    'raw_text': f"ERROR: {e}",
                })

        all_data[class_name] = class_data
    return all_data


def write_excel(all_data, output_path):
    """Write clean, readable Excel with 3 class sheets + raw text reference."""
    wb = Workbook()

    # Styles
    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(wrap_text=True, horizontal='center', vertical='center')
    cell_align = Alignment(wrap_text=True, vertical='top')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    marker_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    marker_font = Font(bold=True, italic=True, size=10, color="548235")

    col_widths = [6, 15, 18, 35, 18, 22, 32, 12, 14]

    for sheet_idx, (class_name, class_data) in enumerate(all_data.items()):
        ws = wb.active if sheet_idx == 0 else wb.create_sheet()
        ws.title = class_name

        # Title row
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
        tc = ws.cell(row=1, column=1,
                     value=f"Khung năng lực và đặc tả kĩ thuật - {class_name}")
        tc.font = title_font
        tc.alignment = Alignment(horizontal='center')

        # Headers
        for ci, hdr in enumerate(HEADERS):
            c = ws.cell(row=2, column=ci + 1, value=hdr)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_align
            c.border = thin_border

        for i, w in enumerate(col_widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w

        row_num = 3
        for img_data in class_data:
            img_num = img_data['image']
            table_rows = img_data['table_rows']
            raw_text = img_data['raw_text']

            # Image marker
            ws.merge_cells(start_row=row_num, start_column=1,
                           end_row=row_num, end_column=9)
            m = ws.cell(row=row_num, column=1,
                        value=f"--- Hình {img_num}.jpg ---")
            m.font = marker_font
            m.fill = marker_fill
            m.alignment = Alignment(horizontal='center')
            row_num += 1

            if table_rows:
                for tr in table_rows:
                    while len(tr) < 9:
                        tr.append('')
                    for ci in range(9):
                        val = str(tr[ci])[:MAX_CELL_TEXT] if ci < len(tr) else ''
                        c = ws.cell(row=row_num, column=ci + 1, value=val)
                        c.alignment = cell_align
                        c.border = thin_border
                    row_num += 1
            else:
                ws.merge_cells(start_row=row_num, start_column=1,
                               end_row=row_num, end_column=9)
                c = ws.cell(row=row_num, column=1,
                            value=raw_text[:MAX_CELL_TEXT] if raw_text
                            else "(Không trích xuất được)")
                c.alignment = cell_align
                row_num += 1

            row_num += 1  # blank separator

        ws.freeze_panes = 'A3'

    # Raw-text reference sheet
    ws_raw = wb.create_sheet(title="Raw Text (Tham khảo)")
    for ci, (hdr, w) in enumerate(
        [("Hình", 8), ("Lớp", 10), ("Nội dung OCR (chỉ chữ in)", 120)], 1
    ):
        ws_raw.cell(row=1, column=ci, value=hdr).font = Font(bold=True, size=11)
        ws_raw.column_dimensions[get_column_letter(ci)].width = w

    raw_row = 2
    for class_name, class_data in all_data.items():
        for img_data in class_data:
            ws_raw.cell(row=raw_row, column=1, value=img_data['image'])
            ws_raw.cell(row=raw_row, column=2, value=class_name)
            ws_raw.cell(row=raw_row, column=3,
                        value=img_data['raw_text'][:MAX_RAW_TEXT])
            ws_raw.cell(row=raw_row, column=3).alignment = cell_align
            raw_row += 1

    wb.save(output_path)
    print(f"\n{'='*60}")
    print(f"✓ Excel saved: {output_path}")
    print(f"{'='*60}")


if __name__ == "__main__":
    output_path = os.path.join(IMG_DIR, "extracted_data.xlsx")

    print("Starting OCR extraction (printed text only)...")
    print(f"Output: {output_path}\n")

    all_data = process_all_images()
    write_excel(all_data, output_path)

    # Summary
    print("\n--- Summary ---")
    for class_name, class_data in all_data.items():
        total_rows = sum(len(d['table_rows']) for d in class_data)
        print(f"  {class_name}: {len(class_data)} images, {total_rows} table rows")
